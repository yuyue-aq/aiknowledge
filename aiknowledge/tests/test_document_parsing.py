from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from pypdf import PdfWriter

from app.domain.documents import DocumentFailureCode, DocumentFormat
from app.services.document_parsing import DocumentParseError, DocumentParser


def test_markdown_parser_preserves_heading_paths_and_text() -> None:
    parser = DocumentParser()

    document = parser.parse(
        filename="产品说明.md",
        content=(
            "# 产品总览\n\n"
            "知潮支持私密与公开知识空间。\n\n"
            "## 上传资料\n\n"
            "支持 PDF、DOCX、Markdown 和 TXT。"
        ).encode(),
    )

    assert document.format is DocumentFormat.MARKDOWN
    assert [block.heading_path for block in document.blocks] == [
        ("产品总览",),
        ("产品总览", "上传资料"),
    ]
    assert document.blocks[1].text == "支持 PDF、DOCX、Markdown 和 TXT。"


def test_txt_parser_detects_a_common_chinese_encoding() -> None:
    parser = DocumentParser()

    document = parser.parse(filename="说明.txt", content="知识库可以追溯来源".encode("gbk"))

    assert document.format is DocumentFormat.TEXT
    assert document.blocks[0].text == "知识库可以追溯来源"


def test_docx_parser_extracts_headings_paragraphs_and_table_text() -> None:
    source = Document()
    source.add_heading("资料说明", level=1)
    source.add_paragraph("这是正文内容。")
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "格式"
    table.cell(1, 1).text = "PDF"
    buffer = BytesIO()
    source.save(buffer)

    document = DocumentParser().parse(filename="资料.docx", content=buffer.getvalue())

    assert document.format is DocumentFormat.DOCX
    assert document.blocks[0].heading_path == ("资料说明",)
    assert document.blocks[0].text == "这是正文内容。"
    assert document.blocks[1].text == "字段 | 值\n格式 | PDF"


def test_pdf_parser_returns_a_specific_failure_for_encrypted_documents() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.encrypt("not-provided")
    buffer = BytesIO()
    writer.write(buffer)

    with pytest.raises(DocumentParseError) as error:
        DocumentParser().parse(filename="加密.pdf", content=buffer.getvalue())

    assert error.value.code is DocumentFailureCode.ENCRYPTED_PDF


def test_pdf_parser_rejects_a_pdf_without_extractable_text() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buffer = BytesIO()
    writer.write(buffer)

    with pytest.raises(DocumentParseError) as error:
        DocumentParser().parse(filename="扫描件.pdf", content=buffer.getvalue())

    assert error.value.code is DocumentFailureCode.EMPTY_DOCUMENT


def test_parser_rejects_malformed_xlsx_before_parsing() -> None:
    with pytest.raises(DocumentParseError) as error:
        DocumentParser().parse(filename="资料.xlsx", content=b"not a spreadsheet")

    assert error.value.code is DocumentFailureCode.MALFORMED_DOCUMENT


def test_csv_parser_turns_faq_rows_into_separate_knowledge_blocks() -> None:
    document = DocumentParser().parse(
        filename="faq.csv",
        content="问题,答案\n如何登录?,使用邮箱和密码登录。\n如何反馈?,在回答下方提交反馈。\n".encode(),
    )
    assert document.format is DocumentFormat.TABLE
    assert len(document.blocks) == 2
    assert document.blocks[0].heading_path == ("FAQ",)
    assert "问题：如何登录?" in document.blocks[0].text
    assert "答案：使用邮箱和密码登录。" in document.blocks[0].text


def test_pptx_parser_extracts_slide_text_and_keeps_slide_numbers() -> None:
    slide_xml = (
        '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<p:cSld><p:spTree><a:t>标题</a:t><a:t>正文内容</a:t></p:spTree></p:cSld></p:sld>'
    ).encode()
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("ppt/slides/slide1.xml", slide_xml)
    document = DocumentParser().parse(filename="演示.pptx", content=buffer.getvalue())
    assert document.format is DocumentFormat.PRESENTATION
    assert document.blocks[0].page_number == 1
    assert document.blocks[0].text == "标题\n正文内容"
