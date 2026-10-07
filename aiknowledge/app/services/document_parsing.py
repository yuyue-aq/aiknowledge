from __future__ import annotations

import csv
from io import BytesIO, StringIO
import posixpath
from pathlib import PurePath
import re
from zipfile import BadZipFile, ZipFile, is_zipfile
from xml.etree import ElementTree

from charset_normalizer import from_bytes
from docx import Document
from markdown_it import MarkdownIt
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from app.services.section_context import page_heading

from app.domain.documents import (
    DocumentBlock,
    DocumentFailureCode,
    DocumentFormat,
    ParsedDocument,
)


class DocumentParseError(ValueError):
    """A stable, user-safe document parsing failure."""

    def __init__(self, code: DocumentFailureCode, message: str) -> None:
        super().__init__(message)
        self.code = code


class DocumentParser:
    """Extract structured text from supported text and office formats."""

    _FORMAT_BY_EXTENSION = {
        ".pdf": DocumentFormat.PDF,
        ".docx": DocumentFormat.DOCX,
        ".md": DocumentFormat.MARKDOWN,
        ".markdown": DocumentFormat.MARKDOWN,
        ".txt": DocumentFormat.TEXT,
        ".csv": DocumentFormat.TABLE,
        ".tsv": DocumentFormat.TABLE,
        ".xlsx": DocumentFormat.TABLE,
        ".pptx": DocumentFormat.PRESENTATION,
    }

    def parse(self, *, filename: str, content: bytes) -> ParsedDocument:
        if not content:
            raise DocumentParseError(
                DocumentFailureCode.EMPTY_DOCUMENT, "文档内容为空，无法处理。"
            )

        document_format = self._resolve_format(filename)
        blocks = {
            DocumentFormat.PDF: self._parse_pdf,
            DocumentFormat.DOCX: self._parse_docx,
            DocumentFormat.MARKDOWN: self._parse_markdown,
            DocumentFormat.TEXT: self._parse_text,
            DocumentFormat.TABLE: lambda value: self._parse_table(value, filename=filename),
            DocumentFormat.PRESENTATION: self._parse_pptx,
        }[document_format](content)

        if not blocks:
            raise DocumentParseError(
                DocumentFailureCode.EMPTY_DOCUMENT,
                "未从文档中提取到可用于知识库的文本。",
            )
        return ParsedDocument(
            filename=filename,
            format=document_format,
            blocks=tuple(blocks),
        )

    def _resolve_format(self, filename: str) -> DocumentFormat:
        extension = PurePath(filename).suffix.lower()
        document_format = self._FORMAT_BY_EXTENSION.get(extension)
        if document_format is None:
            raise DocumentParseError(
                DocumentFailureCode.UNSUPPORTED_FILE_TYPE,
                "支持 PDF、DOCX、Markdown、TXT、CSV/TSV、XLSX 和 PPTX 文件。",
            )
        return document_format

    def _parse_pdf(self, content: bytes) -> list[DocumentBlock]:
        try:
            reader = PdfReader(BytesIO(content))
        except PdfReadError as exc:
            raise DocumentParseError(
                DocumentFailureCode.MALFORMED_DOCUMENT, "PDF 文件损坏或无法读取。"
            ) from exc
        if reader.is_encrypted:
            raise DocumentParseError(
                DocumentFailureCode.ENCRYPTED_PDF, "暂不支持加密 PDF，请移除密码后重新上传。"
            )

        blocks: list[DocumentBlock] = []
        has_image = False
        try:
            for page_number, page in enumerate(reader.pages, 1):
                has_image = has_image or _page_contains_image(page)
                text = _clean_text(page.extract_text() or "")
                if text:
                    blocks.append(
                        DocumentBlock(
                            text=text,
                            heading_path=page_heading(text),
                            ordinal=len(blocks) + 1,
                            page_number=page_number,
                        )
                    )
        except (PdfReadError, ValueError) as exc:
            raise DocumentParseError(
                DocumentFailureCode.MALFORMED_DOCUMENT, "PDF 文本提取失败。"
            ) from exc
        if not blocks and has_image:
            raise DocumentParseError(
                DocumentFailureCode.SCANNED_DOCUMENT,
                "PDF 只包含扫描图像，当前版本暂不支持 OCR。请上传可复制文本的 PDF。",
            )
        return blocks

    def _parse_docx(self, content: bytes) -> list[DocumentBlock]:
        try:
            source = Document(BytesIO(content))
        except Exception as exc:
            raise DocumentParseError(
                DocumentFailureCode.MALFORMED_DOCUMENT, "DOCX 文件损坏或无法读取。"
            ) from exc

        blocks: list[DocumentBlock] = []
        heading_path: list[str] = []
        for paragraph in source.paragraphs:
            text = _clean_text(paragraph.text)
            if not text:
                continue
            level = _heading_level(paragraph.style.name if paragraph.style else "")
            if level is not None:
                heading_path = heading_path[: level - 1]
                heading_path.append(text)
                continue
            blocks.append(
                DocumentBlock(
                    text=text,
                    heading_path=tuple(heading_path),
                    ordinal=len(blocks) + 1,
                )
            )

        for table in source.tables:
            rows = [
                " | ".join(_clean_inline_text(cell.text) for cell in row.cells)
                for row in table.rows
            ]
            table_text = _clean_text("\n".join(row for row in rows if row.strip()))
            if table_text:
                blocks.append(
                    DocumentBlock(
                        text=table_text,
                        heading_path=tuple(heading_path),
                        ordinal=len(blocks) + 1,
                    )
                )
        return blocks

    def _parse_markdown(self, content: bytes) -> list[DocumentBlock]:
        text = self._decode_text(content)
        tokens = MarkdownIt("commonmark").parse(text)
        blocks: list[DocumentBlock] = []
        heading_path: list[str] = []
        index = 0
        while index < len(tokens):
            token = tokens[index]
            if token.type == "heading_open" and index + 1 < len(tokens):
                inline = tokens[index + 1]
                level = int(token.tag[1:])
                heading = _clean_inline_text(inline.content)
                if heading:
                    heading_path = heading_path[: level - 1]
                    heading_path.append(heading)
                index += 3
                continue
            if token.type == "paragraph_open" and index + 1 < len(tokens):
                inline = tokens[index + 1]
                paragraph = _clean_text(inline.content)
                if paragraph:
                    blocks.append(
                        DocumentBlock(
                            text=paragraph,
                            heading_path=tuple(heading_path),
                            ordinal=len(blocks) + 1,
                        )
                    )
                index += 3
                continue
            index += 1
        return blocks

    def _parse_text(self, content: bytes) -> list[DocumentBlock]:
        text = _clean_text(self._decode_text(content))
        return [DocumentBlock(text=text, heading_path=(), ordinal=1)] if text else []

    def _parse_table(self, content: bytes, *, filename: str) -> list[DocumentBlock]:
        if PurePath(filename).suffix.lower() == ".xlsx":
            return self._parse_xlsx(content)
        text = self._decode_text(content)
        delimiter = "\t" if PurePath(filename).suffix.lower() == ".tsv" else ","
        try:
            rows = [
                [cell.strip() for cell in row]
                for row in csv.reader(StringIO(text), delimiter=delimiter)
                if any(cell.strip() for cell in row)
            ]
        except csv.Error as exc:
            raise DocumentParseError(
                DocumentFailureCode.MALFORMED_DOCUMENT, "FAQ 表格格式无法读取。"
            ) from exc
        if not rows:
            return []
        headers = [cell for cell in rows[0]]
        question_index = next(
            (index for index, value in enumerate(headers) if value in {"问题", "问题描述", "question", "Question"}),
            None,
        )
        answer_index = next(
            (index for index, value in enumerate(headers) if value in {"答案", "answer", "Answer"}),
            None,
        )
        has_header = question_index is not None and answer_index is not None
        if not has_header and len(headers) >= 2:
            question_index, answer_index = 0, 1
        return self._table_rows_to_blocks(rows, has_header=has_header, question_index=question_index, answer_index=answer_index)

    @staticmethod
    def _table_rows_to_blocks(
        rows: list[list[str]],
        *,
        has_header: bool,
        question_index: int | None,
        answer_index: int | None,
    ) -> list[DocumentBlock]:
        blocks: list[DocumentBlock] = []
        data_rows = rows[1:] if has_header else rows
        for row in data_rows:
            if question_index is None or answer_index is None:
                line = " | ".join(cell for cell in row if cell)
            else:
                question = row[question_index] if question_index < len(row) else ""
                answer = row[answer_index] if answer_index < len(row) else ""
                line = f"问题：{question}\n答案：{answer}" if question or answer else ""
            line = _clean_text(line)
            if line:
                blocks.append(
                    DocumentBlock(
                        text=line,
                        heading_path=("FAQ",),
                        ordinal=len(blocks) + 1,
                    )
                )
        return blocks

    def _parse_xlsx(self, content: bytes) -> list[DocumentBlock]:
        """Read simple XLSX worksheets without a heavyweight spreadsheet runtime."""

        if not is_zipfile(BytesIO(content)):
            raise DocumentParseError(
                DocumentFailureCode.MALFORMED_DOCUMENT, "XLSX 文件损坏或无法读取。"
            )
        try:
            with ZipFile(BytesIO(content)) as archive:
                names = set(archive.namelist())
                if "xl/workbook.xml" not in names:
                    raise DocumentParseError(
                        DocumentFailureCode.MALFORMED_DOCUMENT, "XLSX 文件缺少工作簿信息。"
                    )
                shared_strings: list[str] = []
                if "xl/sharedStrings.xml" in names:
                    root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
                    for item in root:
                        shared_strings.append(
                            _clean_inline_text("".join(node.text or "" for node in item.iter() if node.tag.endswith("}t")))
                        )
                workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
                relationships: dict[str, str] = {}
                rels_name = "xl/_rels/workbook.xml.rels"
                if rels_name in names:
                    rels_root = ElementTree.fromstring(archive.read(rels_name))
                    for relationship in rels_root:
                        rel_id = relationship.attrib.get("Id")
                        target = relationship.attrib.get("Target")
                        if rel_id and target:
                            relationships[rel_id] = posixpath.normpath(posixpath.join("xl", target))
                rows: list[list[str]] = []
                for sheet in workbook.iter():
                    if not sheet.tag.endswith("}sheet"):
                        continue
                    rel_id = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                    worksheet_name = relationships.get(rel_id or "")
                    if not worksheet_name or worksheet_name not in names:
                        continue
                    root = ElementTree.fromstring(archive.read(worksheet_name))
                    for row in root.iter():
                        if not row.tag.endswith("}row"):
                            continue
                        values: dict[int, str] = {}
                        for cell in row:
                            if not cell.tag.endswith("}c"):
                                continue
                            reference = cell.attrib.get("r", "")
                            column = _xlsx_column_index(reference)
                            if column is None:
                                continue
                            kind = cell.attrib.get("t")
                            value_node = next((node for node in cell if node.tag.endswith("}v")), None)
                            inline_node = next((node for node in cell if node.tag.endswith("}is")), None)
                            if kind == "inlineStr" and inline_node is not None:
                                value = "".join(node.text or "" for node in inline_node.iter() if node.tag.endswith("}t"))
                            elif value_node is not None:
                                value = value_node.text or ""
                                if kind == "s":
                                    try:
                                        value = shared_strings[int(value)]
                                    except (ValueError, IndexError):
                                        value = ""
                            else:
                                value = ""
                            values[column] = _clean_inline_text(value)
                        if values:
                            rows.append([values.get(index, "") for index in range(max(values) + 1)])
                if not rows:
                    return []
                headers = rows[0]
                question_index = next((index for index, value in enumerate(headers) if value in {"问题", "问题描述", "question", "Question"}), None)
                answer_index = next((index for index, value in enumerate(headers) if value in {"答案", "answer", "Answer"}), None)
                has_header = question_index is not None and answer_index is not None
                if not has_header and len(headers) >= 2:
                    question_index, answer_index = 0, 1
                return self._table_rows_to_blocks(
                    rows,
                    has_header=has_header,
                    question_index=question_index,
                    answer_index=answer_index,
                )
        except DocumentParseError:
            raise
        except (BadZipFile, ElementTree.ParseError, KeyError, ValueError, IndexError) as exc:
            raise DocumentParseError(
                DocumentFailureCode.MALFORMED_DOCUMENT, "XLSX 文件损坏或无法读取。"
            ) from exc

    def _parse_pptx(self, content: bytes) -> list[DocumentBlock]:
        try:
            with ZipFile(BytesIO(content)) as archive:
                slide_names = sorted(
                    name
                    for name in archive.namelist()
                    if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
                )
                if not slide_names:
                    raise DocumentParseError(
                        DocumentFailureCode.MALFORMED_DOCUMENT,
                        "PPTX 文件中没有可读取的幻灯片。",
                    )
                blocks: list[DocumentBlock] = []
                for slide_number, name in enumerate(slide_names, 1):
                    root = ElementTree.fromstring(archive.read(name))
                    parts = [
                        _clean_inline_text(element.text or "")
                        for element in root.iter()
                        if element.tag.endswith("}t") and _clean_inline_text(element.text or "")
                    ]
                    text = _clean_text("\n".join(parts))
                    if text:
                        blocks.append(
                            DocumentBlock(
                                text=text,
                                heading_path=(f"幻灯片 {slide_number}",),
                                ordinal=len(blocks) + 1,
                                page_number=slide_number,
                            )
                        )
                return blocks
        except DocumentParseError:
            raise
        except (BadZipFile, ElementTree.ParseError, KeyError) as exc:
            raise DocumentParseError(
                DocumentFailureCode.MALFORMED_DOCUMENT, "PPTX 文件损坏或无法读取。"
            ) from exc

    @staticmethod
    def _decode_text(content: bytes) -> str:
        # Prefer deterministic codecs for our supported Chinese user base. The
        # library fallback covers less common encodings without exposing guesses.
        for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                pass
        match = from_bytes(content).best()
        if match is None:
            raise DocumentParseError(
                DocumentFailureCode.INVALID_TEXT_ENCODING,
                "无法识别 TXT 或 Markdown 文件的文本编码。",
            )
        return str(match)


def _clean_text(text: str) -> str:
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def _clean_inline_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _heading_level(style_name: str) -> int | None:
    match = re.search(r"(?:heading|标题)\s*(\d+)", style_name, re.IGNORECASE)
    return int(match.group(1)) if match else None


def _page_contains_image(page: object) -> bool:
    """Return whether a PDF page has an image XObject (usually a scan)."""

    try:
        resources = page.get("/Resources")  # type: ignore[union-attr]
        if resources is None:
            return False
        resources = resources.get_object() if hasattr(resources, "get_object") else resources
        xobjects = resources.get("/XObject") if hasattr(resources, "get") else None
        if xobjects is None:
            return False
        xobjects = xobjects.get_object() if hasattr(xobjects, "get_object") else xobjects
        for item in xobjects.values():
            resolved = item.get_object() if hasattr(item, "get_object") else item
            if resolved.get("/Subtype") == "/Image":
                return True
    except (AttributeError, KeyError, TypeError, ValueError):
        return False
    return False


def _xlsx_column_index(reference: str) -> int | None:
    match = re.match(r"([A-Za-z]+)", reference)
    if match is None:
        return None
    value = 0
    for char in match.group(1).upper():
        value = value * 26 + (ord(char) - ord("A") + 1)
    return value - 1
